from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path
import json

BASE = Path(__file__).resolve().parent
RESULTS = json.loads((BASE / 'simulation_results.json').read_text())

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.68)
sec.bottom_margin = Inches(0.68)
sec.left_margin = Inches(0.78)
sec.right_margin = Inches(0.78)

# Palette and base typography
styles = doc.styles
styles['Normal'].font.name = 'Aptos'
styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'), 'Aptos')
styles['Normal'].font.size = Pt(9.5)
styles['Normal'].font.color.rgb = RGBColor(38, 48, 58)
styles['Normal'].paragraph_format.space_after = Pt(5)
styles['Normal'].paragraph_format.line_spacing = 1.08
for style_name, size, color in [('Title', 25, '14324A'), ('Heading 1', 16, '14324A'), ('Heading 2', 12, '176B87'), ('Heading 3', 10.5, '176B87')]:
    st = styles[style_name]
    st.font.name = 'Aptos Display' if style_name in ['Title','Heading 1','Heading 2'] else 'Aptos'
    st._element.rPr.rFonts.set(qn('w:eastAsia'), st.font.name)
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = RGBColor.from_string(color)
    st.paragraph_format.space_before = Pt(9 if style_name != 'Title' else 0)
    st.paragraph_format.space_after = Pt(4)
    st.paragraph_format.keep_with_next = True

# Header/footer
header = sec.header.paragraphs[0]
header.text = 'CRICKET BIOMECHANICS  |  STOCHASTIC RECOVERY MODEL'
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
for r in header.runs:
    r.font.name = 'Aptos'; r.font.size = Pt(7); r.font.bold = True; r.font.color.rgb = RGBColor(91, 112, 126)
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
rr = footer.add_run('Preprint-style technical model • Version 1.0 • Illustrative parameters, not athlete-validated')
rr.font.name = 'Aptos'; rr.font.size = Pt(7); rr.font.color.rgb = RGBColor(91, 112, 126)


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), fill); tcPr.append(shd)


def set_cell_text(cell, text, bold=False, color='26303A', size=8.5):
    cell.text = ''
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run(str(text)); r.bold = bold; r.font.name = 'Aptos'; r.font.size = Pt(size); r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def table(headers, rows, widths=None, font=8.2):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = 'Table Grid'
    for j, h in enumerate(headers):
        set_cell_text(t.rows[0].cells[j], h, True, 'FFFFFF', font)
        shade(t.rows[0].cells[j], '14324A')
    for i, row in enumerate(rows):
        cells = t.add_row().cells
        for j, val in enumerate(row):
            set_cell_text(cells[j], val, False, '26303A', font)
            if i % 2 == 1:
                shade(cells[j], 'F1F5F7')
    if widths:
        for row in t.rows:
            for j, w in enumerate(widths):
                row.cells[j].width = Inches(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return t


def para(text='', bold_prefix=None, italic=False, align=None, style=None):
    p = doc.add_paragraph(style=style)
    if align is not None: p.alignment = align
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix); r.bold = True
        p.add_run(text[len(bold_prefix):])
    else:
        r = p.add_run(text); r.italic = italic
    return p


def equation(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(text)
    r.font.name = 'Cambria Math'; r.font.size = Pt(10.5); r.font.bold = True; r.font.color.rgb = RGBColor(23, 107, 135)
    return p


def fig(path, caption, width=6.25):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2); p.paragraph_format.space_after = Pt(1)
    p.add_run().add_picture(str(BASE / path), width=Inches(width))
    cp = doc.add_paragraph(); cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.space_after = Pt(6)
    r = cp.add_run(caption); r.italic = True; r.font.size = Pt(8); r.font.color.rgb = RGBColor(81, 97, 108)

# Title
p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(18); p.paragraph_format.space_after = Pt(5)
r = p.add_run('A Kinematic and Probabilistic Model of Foot Recovery During Cricket Stumping')
r.font.name = 'Aptos Display'; r.font.size = Pt(24); r.font.bold = True; r.font.color.rgb = RGBColor(20,50,74)
sub = doc.add_paragraph(); sub.paragraph_format.space_after = Pt(12)
r = sub.add_run('Sliding versus lifting, the deadline trade-off, and an optimal-clearance calculation')
r.font.name = 'Aptos Display'; r.font.size = Pt(13); r.font.color.rgb = RGBColor(23,107,135)
meta = doc.add_paragraph(); meta.paragraph_format.space_after = Pt(11)
r = meta.add_run('Technical preprint-style study  |  Model version 1.0  |  11 October 2026')
r.font.size = Pt(8.5); r.font.color.rgb = RGBColor(91,112,126)

doc.add_heading('Abstract', level=1)
para('A batter who is outside the ground during a potential stumping faces a constrained motor-control problem: some eligible part of the person or bat must be grounded behind the popping crease at the instant the wicket is put down. This paper develops a transparent, two-dimensional model of one-foot recovery and compares a ground-contact slide with a lifted-and-planted return. Smooth endpoint-constrained trajectories yield closed-form movement-time expressions. A probabilistic extension accounts for uncertainty in the available movement deadline and in the minimum foot clearance required to avoid catching the surface. In an illustrative baseline with a 0.25 m return distance, the slide model requires 369 ms, while the success-probability-maximising lift uses a 4.29 cm peak clearance and requires 319 ms. Estimated success probability is 0.702 for the lift model versus 0.439 for the slide model under the specific assumptions used. These numbers are demonstrations, not empirical estimates of elite batters: the acceleration limits and uncertainty distributions are not calibrated from stumping trials. The main result is structural: if clearance is deterministic and only movement time matters, the fastest lift uses the smallest safe clearance; a non-zero optimum emerges only when added height reduces a modeled collision risk. We provide reproducible Python code, plots, and a protocol for empirical validation.')

doc.add_heading('Keywords', level=2)
para('Cricket; stumping; biomechanics; foot recovery; trajectory optimisation; chance-constrained control; Monte Carlo simulation.')

doc.add_heading('1. Introduction and research question', level=1)
para('A stumping attempt can be framed as a race between the batter re-establishing a legally sufficient ground contact and the wicketkeeper fairly putting down the wicket. The movement choice considered here is narrower: does a batter recover a foot by sliding/dragging it along the surface, or by lifting it, moving it back, and planting it behind the crease? The slide may avoid a vertical excursion but must overcome an effective ground-contact resistance. The lifted path reduces surface interaction but adds vertical travel and demands accurate re-contact.')
para('The formal objective is to find a movement strategy and peak clearance h that maximise the probability of safe, grounded contact before the wicket-breaking deadline. Computer models and motion-capture studies of swing-leg biomechanics demonstrate that leg motion can be modelled through segment kinematics and optimisation, although the cited literature does not directly study stumping recovery [3,4].')
para('Research questions: (i) what movement-time equations follow from explicit trajectory and acceleration constraints? (ii) when is lift faster than slide? (iii) when does uncertainty create a non-zero optimal lift height? (iv) what measurements are needed to make the result player-specific?')

doc.add_heading('2. Legal endpoint and model boundary', level=1)
para('Under the MCC Laws, a batter is out of their ground unless some part of the person or bat is grounded behind the popping crease; stumping has additional conditions, including that the striker is not attempting a run and the delivery is not called No ball [1,2]. The model therefore defines success at the instant a legally eligible contact point is grounded behind the crease, not merely when the toe passes across the line while airborne.')
para('For tractability, only one foot is treated as the eligible contact point. The model omits alternative contacts through the bat, hand, knee, or other body part; it also omits the batter’s trunk motion, joint-level torques, shoe geometry, the keeper’s hand/bail dynamics, and any Law 30.1.2 continuity-of-forward-movement exception. Accordingly, it is a focused foot-recovery model, not a complete law adjudication or a full-body digital twin.')

doc.add_heading('3. Kinematic model and derivation', level=1)
doc.add_heading('3.1 State variables and smooth endpoint constraints', level=2)
para('Let d > 0 be the horizontal distance from the foot’s current effective contact point to a safe landing point behind the crease (including a small margin). Let T be movement duration, h the peak vertical clearance, and u = t/T in [0,1]. Movement begins and ends at zero horizontal velocity so that the foot is not required to teleport or stop instantaneously.')
equation('x(t) = d(3u^2 - 2u^3),     u = t/T')
para('Differentiating with respect to t gives x-dot = (6d/T)u(1-u) and x-double-dot = (6d/T^2)(1-2u). Therefore max |x-double-dot| = 6d/T^2. This is a deliberately simple smooth-start/smooth-stop trajectory; it is not asserted to be the unique or human-optimal motor pattern.')

doc.add_heading('3.2 Slide/drag strategy', level=2)
para('Let a_m be an effective maximum horizontal acceleration available to move the foot and let a_r be an effective acceleration penalty representing sliding resistance and other surface-contact losses. The kinematic acceleration demand is bounded by the remaining acceleration budget:')
equation('6d/T_slide^2 + a_r <= a_m')
para('If a_m > a_r, the minimum time under this simplified bound is')
equation('T_slide = sqrt[6d / (a_m - a_r)]')
para('If a_m <= a_r, the model declares this slide trajectory infeasible. The parameter a_r is not a claim that a shoe follows ideal Coulomb sliding friction; it is a lumped, experimentally calibratable penalty. A more detailed model would identify friction, shoe-ground contact mode, and joint moments separately.')

doc.add_heading('3.3 Lift-and-plant strategy', level=2)
para('For the lifted strategy, use the same horizontal trajectory and a smooth vertical arch with zero vertical velocity at take-off and landing:')
equation('z(t) = 16h u^2(1-u)^2')
para('This trajectory has z(0) = z(T) = 0 and z(T/2) = h. Differentiation gives z-double-dot = (32h/T^2)(1 - 6u + 6u^2), whose maximum absolute magnitude is 32h/T^2. Under a task-space acceleration-norm bound a_l, and using the endpoint where the horizontal and vertical acceleration magnitudes are jointly largest,')
equation('sqrt[(6d/T^2)^2 + (32h/T^2)^2] <= a_l')
para('Rearranging yields the movement-time lower bound / trajectory time used in the simulation:')
equation('T_lift(h) = sqrt{ sqrt(36d^2 + 1024h^2) / a_l }')
para('The nested square root is intentional: the inner expression has units of metres; division by acceleration gives seconds squared, and the outer square root gives seconds. The model treats the acceleration norm as an effective movement-capacity parameter rather than a direct measurement of a specific muscle or joint.')

doc.add_heading('3.4 Slide-versus-lift crossover', level=2)
para('For a specified height h, lifting is faster if T_lift(h) < T_slide. Equating the two times gives the clearance at which movement durations are equal:')
equation('h_cross = sqrt{[(a_l T_slide^2)^2 - 36d^2] / 1024}')
para('This real-valued crossover exists only when the numerator is positive. If no such non-negative value exists, the assumed lift acceleration is not enough to make the lift trajectory faster at any height in this model. For the baseline parameters, h_cross = 7.12 cm: below that height the model’s lift trajectory is faster; above it the extra vertical travel makes it slower than sliding.')

doc.add_heading('3.5 Why a “best height” needs a risk term', level=2)
para('Since T_lift(h) increases with h for h > 0, a time-only optimisation always chooses the smallest feasible clearance. There is no interior height optimum in that deterministic model. To capture the question’s trade-off, we add a terrain/foot-trajectory uncertainty: too little clearance can catch the pitch or fail to clear a local irregularity, while too much clearance costs time.')
para('Let Delta be the movement time available after the recovery movement starts, and let H be the minimum clearance needed on that trial. Assume, for an initial tractable model, independent normal distributions Delta ~ N(mu_Delta, sigma_Delta^2) and H ~ N(mu_H, sigma_H^2). Then the lifted strategy’s success probability is')
equation('P_lift(h) = Phi[(mu_Delta - T_lift(h))/sigma_Delta] * Phi[(h - mu_H)/sigma_H]')
para('Here Phi is the standard normal cumulative distribution function. The first factor is the chance of completing the trajectory before the randomly varying deadline; the second is the chance that peak clearance exceeds the random clearance demand. The independence assumption is a modelling convenience to be tested, not a known property of real movement.')
para('Write z_T = (mu_Delta - T_lift(h))/sigma_Delta, z_H = (h - mu_H)/sigma_H, and let phi be the standard normal density. Differentiating log P_lift and setting it to zero gives the first-order condition:')
equation('[phi(z_H) / (sigma_H Phi(z_H))] = T_lift\'(h) * [phi(z_T) / (sigma_Delta Phi(z_T))]')
para('This is the central trade-off equation: the marginal gain in clearance reliability equals the marginal loss in deadline reliability. With S(h) = sqrt(36d^2 + 1024h^2), the time derivative is')
equation('T_lift\'(h) = 512h / [sqrt(a_l) * S(h)^(3/2)]')
para('There is generally no useful closed-form solution for h because normal-CDF ratios appear in the first-order condition. We therefore maximise P_lift(h) numerically on a height grid and verify the result by Monte Carlo sampling.')

doc.add_heading('4. Simulation design', level=1)
para('The code uses SI units, a fixed random seed, 300,000 Monte Carlo trials, and a one-dimensional grid of 500 candidate heights from 2 mm to 12 cm. At every candidate height, it samples an available movement budget and a minimum required clearance, then counts a success if T_lift(h) <= Delta and h >= H. The same sampled deadlines are used to estimate slide timing success. Results are therefore reproducible up to numerical-library differences.')
table(['Parameter','Baseline','Interpretation'], [
    ('d','0.25 m','Required horizontal return to a safe grounded point'),
    ('a_m','15 m/s^2','Effective slide movement acceleration limit'),
    ('a_r','4 m/s^2','Effective contact-resistance penalty during sliding'),
    ('a_l','20 m/s^2','Effective task-space acceleration bound for lift'),
    ('mu_Delta','0.360 s','Mean movement time available after movement starts'),
    ('sigma_Delta','0.060 s','Standard deviation of available movement time'),
    ('mu_H','0.025 m','Mean clearance demand from surface/trajectory uncertainty'),
    ('sigma_H','0.012 m','Standard deviation of clearance demand'),
], widths=[1.15,1.25,4.2], font=8.1)
para('Important parameter caveat: these values are illustrative engineering assumptions selected to make the trade-off visible. They are not estimates from an instrumented cricket dataset and should not be used to prescribe technique to players.')

doc.add_heading('5. Results', level=1)

table(['Outcome','Slide/drag','Lift at optimum'], [
    ('Movement time','369.3 ms','318.9 ms at h = 4.29 cm'),
    ('Timing-only success probability','0.4386','0.753 approx.'),
    ('Clearance-success probability','Not separately modelled (contact maintained)','0.932 approx.'),
    ('Combined success probability','0.4386','0.7024'),
    ('Monte Carlo check','0.4383','0.7021 near optimum'),
], widths=[2.15,2.0,2.45], font=8.1)
para('For the baseline case, numerical maximisation gives h* = 4.291 cm. At this height, the lift movement takes 318.9 ms. The analytic success probability is 0.7024, and the Monte Carlo estimate near the optimum is 0.7021, a difference of 0.0003. For the sliding strategy, movement time is 369.3 ms and the probability of beating the uncertain timing deadline is 0.4386 (Monte Carlo: 0.4383). The slide/lift equal-time clearance is 7.12 cm.')
fig('time_vs_height.png', 'Figure 1. Movement time rises with the peak height of the lifted foot. The horizontal dashed line is the modelled slide time.')
fig('success_probability_vs_height.png', 'Figure 2. A non-zero optimum appears when extra clearance reduces contact risk but also consumes movement time. Dots show Monte Carlo checks.')
fig('strategy_phase_diagram.png', 'Figure 3. Deterministic strategy map over return distance and available movement budget, using the optimised baseline height as the reference lift height.')

doc.add_heading('5.1 Sensitivity to return distance', level=2)
table(['Return distance d','Slide time','Best lift h*','Lift time at h*','Slide timing probability','Lift success probability'], [
    ('0.15 m','286 ms','4.48 cm','291 ms','0.891','0.832'),
    ('0.25 m','369 ms','4.29 cm','319 ms','0.439','0.702'),
    ('0.35 m','437 ms','4.13 cm','352 ms','0.100','0.503'),
], widths=[1.12,0.9,0.95,1.05,1.2,1.1], font=7.7)
para('The strategy ranking changes with distance. At d = 0.15 m, sliding is quicker and has the higher success probability under these assumptions; for longer distances, the modeled lift becomes preferable. This is the expected consequence of the sliding-resistance penalty: its time grows more rapidly with d than the lift model’s time at the selected height.')

doc.add_heading('5.2 Sensitivity to the available-time distribution', level=2)
table(['Mean available time','Optimising lift height','Optimised lift success probability'], [
    ('280 ms','3.53 cm','0.265'),
    ('320 ms','3.89 cm','0.484'),
    ('360 ms','4.29 cm','0.702'),
    ('400 ms','4.72 cm','0.862'),
    ('440 ms','5.21 cm','0.949'),
], widths=[1.75,2.1,2.2], font=8.1)
para('When a longer movement budget is typical, a somewhat higher clearance is affordable and gives additional protection against clearance failure. When the deadline is tighter, the optimum shifts down because the time penalty becomes more costly. Thus a single universal “best height” is not supported even inside this simple model.')

doc.add_heading('6. Interpretation', level=1)
para('The model produces three distinct statements. First, legal success is about being grounded behind the crease at the relevant instant, not merely moving a foot across the line in the air [1,2]. Second, under a deterministic height constraint, the optimal lift is the smallest clearance that reliably avoids surface contact; adding height only increases time in the present trajectory family. Third, once clearance uncertainty is modelled, an interior maximum can emerge because each extra centimetre buys reliability but costs time. In the illustrative case, that balance occurs near 4.3 cm.')
para('The numerical result should not be read as “batters should lift their feet 4.3 cm.” It should be read as: with this trajectory family and these assumptions, the mathematical optimum lies there. Player-specific predictions require measuring the return distance, reaction/movement onset, joint dynamics, surface contact, the keeper’s stump-breaking time distribution, and the competing eligible contact paths. The bat may provide an earlier grounded point and could outperform either foot strategy; that possibility is outside this comparison.')

doc.add_heading('7. Empirical validation plan', level=1)
para('A research-grade follow-up should replace assumed parameters with measurements and test the predicted crossover and optimum. A practical pilot could use the following protocol:')
table(['Stage','Protocol','Measured outcome'], [
    ('1. Instrumented setup','Use a safe practice wicket, marked crease, high-speed cameras at >= 240 fps (preferably 500 fps), and synchronized audio/LED or contact sensors. Add 3D motion capture and a force plate/pressure insole if available.','Foot, bat, and body contact states; keeper glove/bail event; coordinate uncertainty.'),
    ('2. Participants','Recruit batters across skill levels; after a pilot, conduct a power analysis. A starting design might be 20 batters x 20-30 repeated trials, randomised across conditions. Obtain informed consent and ethics approval.','Between-player variability and within-player repeatability.'),
    ('3. Trial conditions','Randomise required return distance (e.g., 0.10-0.40 m), slide vs lift instructions, dry vs slightly varied safe surface conditions, and an externally controlled or recorded wicket-breaking deadline. Avoid creating unsafe live-stumping collisions.','Trajectory time, peak toe clearance, clearance/contact failures, safe-grounding margin.'),
    ('4. Parameter fitting','Fit movement trajectories and acceleration limits from coordinates. Estimate distribution of available movement budgets from synchronized keeper events, and clearance demand from observed near-catches or a safe surface-proxy experiment.','Player- and condition-specific parameters with confidence intervals.'),
    ('5. Validation','Fit on a subset of trials, predict held-out trials, and use player-level cross-validation. Compare predicted versus observed safe-return probabilities and calibration plots.','Prediction error, calibration, and uncertainty intervals.'),
], widths=[1.0,3.95,1.7], font=7.3)
para('A stronger model would use an articulated hip-knee-ankle-foot chain with bounded joint torques, foot-ground complementarity, and the true geometric condition that any legally eligible part of the person or bat is grounded behind the crease. An optimal-control solver could then minimise expected dismissal probability or maximise a safety margin, rather than relying on a stylised task-space acceleration bound. Results should be compared against simple alternatives: plant the bat behind the crease, slide the foot, and lift-and-plant the foot.')

doc.add_heading('8. Limitations', level=1)
para('The study is theoretical and its parameters are not athlete-calibrated. The endpoint trajectory is a convenient cubic/smooth polynomial, not an identified human movement primitive. The acceleration constraints approximate complex joint-level limits; they do not model muscle activation, fatigue, body balance, footwear, pitch deformation, or the need to avoid losing balance. Independence of timing and clearance noise may be false. The sliding model folds ground interaction into one acceleration penalty. The wicketkeeper deadline is represented as a random available-time budget, rather than being generated from measured glove-to-bail kinematics. The model also omits the bat and other body contacts that may satisfy the cricket law sooner than returning a foot.')
para('Consequently, the model supports hypotheses and experimental design, not match outcome prediction. The parameters should be published with confidence intervals once calibrated, and the optimum should be reported as a distribution across players and conditions rather than a universal constant.')

doc.add_heading('9. Conclusion', level=1)
para('A compact mathematical model can express the foot-recovery trade-off. The slide time is T_slide = sqrt[6d/(a_m-a_r)]; the lift time is T_lift(h) = sqrt{sqrt(36d^2+1024h^2)/a_l}. A timing-only lift objective always favours minimum safe clearance. A non-zero best height appears only after incorporating the probability that insufficient clearance causes contact failure, giving P_lift(h) = Phi[(mu_Delta-T_lift(h))/sigma_Delta] Phi[(h-mu_H)/sigma_H]. Numerical optimisation of the illustrative baseline gives h* around 4.3 cm and a modelled success probability of 0.702. The next scientific step is not to treat that value as coaching advice, but to obtain high-speed motion/contact data and validate the model across batters, distances, surfaces, and keeper timing.')

doc.add_heading('References', level=1)
refs = [
    '[1] Marylebone Cricket Club (MCC). Law 39: Stumped. The Laws of Cricket. https://www.lords.org/mcc/the-laws-of-cricket/stumped',
    '[2] Marylebone Cricket Club (MCC). The Laws of Cricket: Law Appendices (including ground, bat and person definitions). https://www.lords.org/mcc/the-laws-of-cricket/law-appendices',
    '[3] Rottier, T. D., & Allen, S. J. (2021). The influence of swing leg technique on maximum running speed. Journal of Biomechanics, 126, 110640. https://doi.org/10.1016/j.jbiomech.2021.110640',
    '[4] Nagahara, R., Matsubayashi, T., Matsuo, A., & Zushi, K. (2017). Alteration of swing leg work and power during human accelerated sprinting. Biology Open, 6(5), 633-641. https://doi.org/10.1242/bio.024281',
]
for ref in refs:
    p = doc.add_paragraph(ref)
    p.paragraph_format.left_indent = Inches(0.15)
    p.paragraph_format.first_line_indent = Inches(-0.15)
    p.paragraph_format.space_after = Pt(4)
    for r in p.runs:
        r.font.size = Pt(8.2)

# Add a small model reproducibility note

doc.add_heading('Reproducibility', level=1)
para('Run `python simulate_stumping.py` with Python 3, NumPy, and Matplotlib. The script writes the three figures and `simulation_results.json` into its own directory. The random seed is fixed to 20261011. All baseline constants are declared at the top of the script and can be changed for sensitivity experiments.')

out = BASE / 'stumping_recovery_research_paper.docx'
doc.save(out)
print(out)
